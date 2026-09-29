"""Shared fixtures, and the isolation that makes the suite safe to run.

Why the first thing in this file is an environment mutation
-----------------------------------------------------------
``app.config`` builds a :class:`~app.config.Settings` instance *at import time*
(``settings = get_settings()``) and ``app.database.base`` creates the SQLAlchemy engine
from it, also at import time. By the time any test module has said ``import app``, the
database URL and every writable directory are already fixed. So the redirection has to
happen before the first ``app.*`` import in the process, which is here, at module scope,
above the imports - not in a fixture.

What that buys is the only property that matters for a suite that ingests files, writes
reports and seeds a control library: **it cannot touch the developer's real data.**
``data/audit.db`` and ``data/uploads`` are never opened. ``tests/test_config.py`` asserts
that rather than trusting it, by checking the live engine URL resolves inside the
temporary tree.

The one directory deliberately *not* redirected is ``controls_dir``. The control library
JSON is a read-only study artefact that the seeding path needs, and copying it would test
a copy rather than the file the application ships.

No network, anywhere
--------------------
``LLM_PROVIDER=mock`` and ``EMBEDDING_PROVIDER=local`` select the two offline
implementations, and both API-key variables are removed from the environment so a
developer's real key cannot leak into a run. ``tests/test_llm_providers.py`` proves the
OpenAI-compatible provider is never constructible here and that the mock is what runs.

Fixture cost
------------
Generating the synthetic evidence files is session-scoped (it writes real .docx/.xlsx/
.pdf bytes); everything touching the database is function-scoped over a freshly created
schema, because a shared database between tests would make ordering matter.
"""

from __future__ import annotations

import os
import shutil
import sys
import tempfile
from pathlib import Path

# --------------------------------------------------------------------------- isolation
# Nothing above this block may import from ``app``.
if "app.config" in sys.modules:  # pragma: no cover - defensive
    raise RuntimeError(
        "app.config was imported before tests/conftest.py could redirect the database and "
        "storage directories. The suite would run against the developer's real data."
    )

TMP_ROOT = Path(tempfile.mkdtemp(prefix="itaudit-tests-")).resolve()

_ENVIRONMENT = {
    "DATABASE_URL": "sqlite:///" + str(TMP_ROOT / "audit-test.db"),
    "DATA_DIR": str(TMP_ROOT / "data"),
    "UPLOAD_DIR": str(TMP_ROOT / "uploads"),
    "REPORT_DIR": str(TMP_ROOT / "reports"),
    "SYNTHETIC_DIR": str(TMP_ROOT / "synthetic"),
    "EVALUATION_OUTPUT_DIR": str(TMP_ROOT / "evaluation"),
    "SAMPLE_EVIDENCE_DIR": str(TMP_ROOT / "sample_evidence"),
    # Offline everything.
    "LLM_PROVIDER": "mock",
    "EMBEDDING_PROVIDER": "local",
    "MOCK_HALLUCINATION_RATE": "0.0",
    "MOCK_SEED": "1337",
    "USE_API": "false",
    "DB_ECHO": "false",
}
os.environ.update(_ENVIRONMENT)

# Removed rather than blanked: a real key present in the environment must not reach a
# provider constructor even by accident, and an empty string is still "set" to some
# libraries.
for _secret in (
    "LLM_API_KEY",
    "LLM_BASE_URL",
    "EMBEDDING_API_KEY",
    "EMBEDDING_BASE_URL",
    "ANTHROPIC_API_KEY",
    "ANTHROPIC_BASE_URL",
):
    os.environ.pop(_secret, None)

import pytest  # noqa: E402

import app.config as app_config  # noqa: E402

# Explicitly clear the cache the module populated on import. The values are identical
# either way - the environment was already redirected - but doing it here means the
# suite depends on the documented API rather than on import order staying as it is.
app_config.get_settings.cache_clear()
SETTINGS = app_config.get_settings()

from app.audit import service as audit_service  # noqa: E402
from app.audit.engine import AssessmentEngine  # noqa: E402
from app.database.base import Base, SessionLocal, engine, init_db  # noqa: E402
from app.database.seed import seed_controls  # noqa: E402
from app.evaluation.datasets import get_dataset  # noqa: E402
from app.evaluation.generator import generate_dataset  # noqa: E402
from app.evidence.service import ingest_file  # noqa: E402
from app.llm.base import LLMError, LLMProvider, LLMResponse  # noqa: E402
from app.llm.mock_provider import MockLLMProvider  # noqa: E402
from app.schemas.enums import EvidenceType, ExperimentMode, HumanDecision  # noqa: E402

# ---- known properties of the hand-built fixture evidence
#
# All four artefacts describe the same fabricated MFA control, so a project holding all
# of them is coherent: the correct conclusion for CONTROL-001 over this evidence is
# POTENTIAL_DEFICIENCY (five of fifty privileged accounts are not enrolled). Mixing in an
# unrelated password-policy artefact would let a finding about *passwords* surface while
# an *MFA* control was under test, which makes every fixture-based assertion harder to
# read than it needs to be.
#: Rows in the fixture privileged-account export, excluding the header.
FIXTURE_CSV_ROWS = 50
#: Accounts in that file whose MFA_Status is Disabled.
FIXTURE_CSV_EXCEPTIONS = 5
#: Their 1-based spreadsheet row numbers (header occupies row 1), so a test can assert
#: that a locator points at the row an auditor would open the file to.
FIXTURE_CSV_EXCEPTION_ROWS = [7, 14, 23, 38, 45]
#: A distinctive sentence written into the fixture policy .txt.
FIXTURE_POLICY_SENTENCE = (
    "Every account classified as privileged must be enrolled in multi-factor authentication "
    "before privileged access is granted."
)
#: The first setting row of the fixture configuration export, as the parser renders it.
FIXTURE_CONFIG_FIRST_ROW = "row 2: MFA required for privileged roles | Enabled"


def pytest_sessionfinish(session, exitstatus):  # noqa: ARG001 - pytest hook signature
    """Remove the temporary tree. Left in place when a test failed, for inspection."""
    if exitstatus == 0:
        shutil.rmtree(TMP_ROOT, ignore_errors=True)


# --------------------------------------------------------------------------- settings
@pytest.fixture(scope="session")
def tmp_root() -> Path:
    """The temporary tree every writable path in this run lives under."""
    return TMP_ROOT


@pytest.fixture(scope="session")
def settings():
    """The live application settings - the same object the application code reads."""
    return SETTINGS


class RemoteClientConstructed(AssertionError):
    """Raised if anything in the suite tries to build a real OpenAI-compatible client."""


@pytest.fixture(scope="session", autouse=True)
def no_remote_llm_client():
    """Make constructing a real provider client an immediate, loud failure.

    Removing the API key from the environment already means the factory returns the
    mock, but that is a fact about configuration. This closes the remaining gap: if any
    code path ever reaches ``openai.OpenAI(...)`` - because a test built a provider by
    hand, or because a fallback stopped working - the suite fails at that line with a
    named error instead of silently attempting a network call.

    Yields the exception type so a test can assert the guard is actually armed.
    """
    import app.llm.openai_provider as openai_provider

    sdk = openai_provider._openai
    if sdk is None:  # pragma: no cover - only on an install without the SDK
        yield RemoteClientConstructed
        return

    original = sdk.OpenAI

    def _refuse(*args, **kwargs):
        raise RemoteClientConstructed(
            "A remote LLM client was constructed during the test suite. The suite must "
            "run entirely offline against the deterministic mock provider."
        )

    sdk.OpenAI = _refuse
    try:
        yield RemoteClientConstructed
    finally:
        sdk.OpenAI = original


# --------------------------------------------------------------------------- database
@pytest.fixture(scope="session", autouse=True)
def _schema() -> None:
    """Create the schema once so the first test does not pay for it."""
    init_db()


@pytest.fixture()
def clean_database() -> None:
    """Drop and recreate every table, so no test can see another test's rows.

    Dropping rather than truncating keeps the fixture honest about schema changes: a
    model that stops being creatable fails here rather than surfacing as a confusing
    query error three tests later.
    """
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)


@pytest.fixture()
def db_session(clean_database):  # noqa: ARG001 - ordering dependency
    """A session on an empty database."""
    session = SessionLocal()
    try:
        yield session
    finally:
        session.close()


@pytest.fixture()
def seeded_session(db_session):
    """A session on a database holding the full synthetic control library."""
    seed_controls(db_session)
    return db_session


@pytest.fixture()
def control(seeded_session):
    """CONTROL-001 - the MFA control three of the six datasets are written against."""
    return audit_service.require_control(seeded_session, "CONTROL-001")


@pytest.fixture()
def project(seeded_session):
    """An audit project with CONTROL-001 in scope and no evidence yet."""
    return audit_service.create_project(
        seeded_session,
        name="Privileged Access Review (test)",
        audit_area="Identity and Access Management",
        description="Synthetic project used by the test suite.",
        auditor_name="Test Auditor",
        control_refs=["CONTROL-001"],
        actor="Test Auditor",
    )


# --------------------------------------------------------------------- evidence files
@pytest.fixture(scope="session")
def sample_files(tmp_root) -> dict:
    """Real evidence files with hand-chosen contents, written once per session.

    These are deliberately *not* the synthetic evaluation datasets: their numbers are
    fixed here so a test can assert against a count it can see in this file, rather than
    against whatever the dataset generator happens to produce.
    """
    directory = tmp_root / "fixture_evidence"
    directory.mkdir(parents=True, exist_ok=True)
    return {
        "csv": _write_exception_csv(directory / "Privileged_Accounts_Export.csv"),
        "policy_txt": _write_policy_txt(directory / "Multi_Factor_Authentication_Policy.txt"),
        "docx": _write_policy_docx(directory / "Access_Control_Policy.docx"),
        "xlsx": _write_config_xlsx(directory / "MFA_Enforcement_Config.xlsx"),
    }


def _write_exception_csv(path: Path) -> Path:
    """A 50-account privileged listing with exactly five MFA exceptions.

    The exceptions are placed by *spreadsheet* row, so data record ``n`` is written at
    row ``n + 1`` (the header occupies row 1). That off-by-one is the whole point of
    ``tests/test_parsers.py``, so it is done explicitly here rather than implied.
    """
    exception_ordinals = {row - 1 for row in FIXTURE_CSV_EXCEPTION_ROWS}
    lines = ["Account,Account_Type,MFA_Status,Last_Login"]
    for ordinal in range(1, FIXTURE_CSV_ROWS + 1):
        status = "Disabled" if ordinal in exception_ordinals else "Enabled"
        lines.append(
            "svc-account-{0:03d},Privileged,{1},2024-06-{2:02d}".format(
                ordinal, status, (ordinal % 28) + 1
            )
        )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def _write_policy_txt(path: Path) -> Path:
    path.write_text(
        "SYNTHETIC RESEARCH DATA - fabricated for a university research prototype.\n"
        "\n"
        "MULTI-FACTOR AUTHENTICATION POLICY (SYNTHETIC)\n"
        "\n"
        "Scope\n"
        "This policy applies to every privileged account in the fabricated environment.\n"
        "\n"
        "Enrolment\n"
        "{0}\n"
        "\n"
        "Audit reporting\n"
        "The privileged account listing produced for audit must report the multi-factor "
        "authentication enrolment status of every account.\n".format(FIXTURE_POLICY_SENTENCE),
        encoding="utf-8",
    )
    return path


def _write_policy_docx(path: Path) -> Path:
    from docx import Document

    document = Document()
    document.add_heading("Access Control Policy (SYNTHETIC)", level=1)
    document.add_paragraph(
        "SYNTHETIC RESEARCH DATA - this document describes no real organisation."
    )
    document.add_heading("Section 2 - Privileged access", level=2)
    document.add_paragraph(
        "All privileged accounts must authenticate with a second factor. The privileged "
        "account listing produced for audit must report the multi-factor authentication "
        "enrolment status of every account."
    )
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Requirement"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "MFA enrolment"
    table.cell(1, 1).text = "Mandatory"
    document.save(str(path))
    return path


def _write_config_xlsx(path: Path) -> Path:
    """A two-sheet identity-provider configuration export, for the multi-sheet tests."""
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "PolicyRules"
    sheet.append(["Setting", "Value", "Data_Origin"])
    sheet.append(["MFA required for privileged roles", "Enabled", "SYNTHETIC-RESEARCH-DATA"])
    sheet.append(["Allowed second factors", "Authenticator app; FIDO2", "SYNTHETIC-RESEARCH-DATA"])
    sheet.append(["Re-authentication interval (hours)", 12, "SYNTHETIC-RESEARCH-DATA"])

    second = workbook.create_sheet("Scope")
    second.append(["Directory", "Privileged_Accounts", "Data_Origin"])
    second.append(["corp.example-synthetic", FIXTURE_CSV_ROWS, "SYNTHETIC-RESEARCH-DATA"])
    workbook.save(str(path))
    return path


@pytest.fixture()
def ingested_project(seeded_session, project, sample_files):
    """``project`` with all four fixture artefacts stored, parsed, chunked and indexed."""
    types = {
        "csv": EvidenceType.USER_LISTING,
        "policy_txt": EvidenceType.POLICY,
        "docx": EvidenceType.POLICY,
        "xlsx": EvidenceType.CONFIGURATION_EXPORT,
    }
    for key, path in sample_files.items():
        ingest_file(
            seeded_session,
            project.id,
            path.read_bytes(),
            path.name,
            evidence_type=types[key],
            description="Fixture evidence for the test suite.",
            uploaded_by="Test Auditor",
            is_synthetic=True,
        )
    return project


# ------------------------------------------------------------- synthetic datasets
@pytest.fixture(scope="session")
def generated_datasets(tmp_root) -> dict:
    """Every synthetic dataset written to disk once, keyed by dataset id.

    Session-scoped because writing .docx, .xlsx and .pdf bytes six times over is the
    single most expensive thing this suite does, and the bytes are deterministic.
    """
    from app.evaluation.datasets import dataset_ids

    return {dataset_id: generate_dataset(dataset_id) for dataset_id in dataset_ids()}


@pytest.fixture()
def dataset_project(seeded_session, generated_datasets):
    """Factory: ``dataset_project("DATASET-005")`` -> ``(dataset, project)``.

    Each dataset gets its own project, exactly as the evaluation harness does it, so
    one dataset's evidence can never be retrieved while assessing another.
    """

    def _build(dataset_id):
        dataset = get_dataset(dataset_id)
        project_row = audit_service.create_project(
            seeded_session,
            name="Evaluation fixture {0}".format(dataset.dataset_id),
            audit_area=dataset.name,
            control_refs=[dataset.control_ref],
            actor="Test Auditor",
        )
        for path in generated_datasets[dataset.dataset_id]:
            declared = dataset.file(path.name)
            ingest_file(
                seeded_session,
                project_row.id,
                path.read_bytes(),
                path.name,
                evidence_type=declared.evidence_type if declared else EvidenceType.OTHER,
                uploaded_by="evaluation-fixture",
                is_synthetic=True,
            )
        return dataset, project_row

    return _build


# ------------------------------------------------------------------------ providers
@pytest.fixture()
def mock_llm(settings) -> MockLLMProvider:
    return MockLLMProvider(settings=settings)


class BrokenLLM(LLMProvider):
    """A provider that always fails, for the "failure is recorded, not swallowed" tests."""

    name = "broken"

    def __init__(self, message: str = "simulated provider outage") -> None:
        super().__init__(model="broken-model")
        self.message = message
        self.calls = 0

    def complete(self, system, user, json_schema=None, temperature=0.0, max_tokens=2000, context=None):
        self.calls += 1
        raise LLMError(self.message)


class ScriptedLLM(LLMProvider):
    """Returns a caller-supplied payload per purpose, so a specific model behaviour
    (a fabricated citation, an overstated conclusion) can be reproduced exactly."""

    name = "scripted"

    def __init__(self, payloads) -> None:
        super().__init__(model="scripted-model")
        self.payloads = dict(payloads)
        self.purposes = []

    def complete(self, system, user, json_schema=None, temperature=0.0, max_tokens=2000, context=None):
        import json

        purpose = (context.purpose if context is not None else "assessment") or "assessment"
        self.purposes.append(purpose)
        payload = self.payloads.get(purpose)
        if payload is None:
            raise LLMError("ScriptedLLM has no payload for purpose {0!r}".format(purpose))
        text = payload if isinstance(payload, str) else json.dumps(payload)
        return LLMResponse(
            text=text,
            provider=self.name,
            model=self.model,
            prompt_tokens=10,
            completion_tokens=20,
            latency_ms=1,
            finish_reason="stop",
        )


@pytest.fixture()
def broken_llm() -> BrokenLLM:
    return BrokenLLM()


@pytest.fixture()
def scripted_llm():
    return ScriptedLLM


@pytest.fixture()
def engine_factory(seeded_session, settings):
    """Factory for an :class:`AssessmentEngine` bound to the test session."""

    def _make(llm=None, retriever=None):
        return AssessmentEngine(seeded_session, llm=llm, retriever=retriever, settings=settings)

    return _make


# ------------------------------------------------------------ assessments / reviews
@pytest.fixture()
def completed_assessment(seeded_session, ingested_project, engine_factory):
    """One persisted mode-C assessment over the fixture evidence."""
    result = engine_factory().assess_control(
        ingested_project.id, "CONTROL-001", mode=ExperimentMode.C_RAG_WORKFLOW
    )
    assert result.assessment_id is not None
    return audit_service.require_assessment(seeded_session, result.assessment_id)


@pytest.fixture()
def completed_review(seeded_session, completed_assessment):
    """An auditor decision accepting that assessment unchanged."""
    return audit_service.record_human_review(
        seeded_session,
        completed_assessment.id,
        reviewer_name="Test Auditor",
        decision=HumanDecision.ACCEPTED,
        comments="Reviewed against source evidence.",
        review_seconds=42.0,
        usefulness_rating=4,
    )


# ------------------------------------------------------------------------------ API
@pytest.fixture()
def api_client(clean_database):  # noqa: ARG001 - ordering dependency
    """A ``TestClient`` over a freshly created schema.

    The app's own lifespan seeds the control library and the demo project, so this
    fixture yields the same starting state a real deployment has on first boot.
    """
    from fastapi.testclient import TestClient

    from app.api.main import app

    with TestClient(app) as client:
        yield client


@pytest.fixture()
def api_session():
    """A direct session for asserting what the API actually wrote."""
    session = SessionLocal()
    try:
        yield session
    finally:
        session.close()
